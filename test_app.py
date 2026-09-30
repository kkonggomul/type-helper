import io
import ast
import re
import unittest
from pathlib import Path
from unittest.mock import patch, MagicMock

from docx import Document
from PIL import Image
from streamlit.testing.v1 import AppTest

from transcription import transcription_to_docx_buffer, TRANSCRIPTION_PROMPT


class ExportTests(unittest.TestCase):
    def test_existing_exam_export(self):
        tree = ast.parse(Path(__file__).with_name('app.py').read_text(encoding='utf-8'))
        functions = [node for node in tree.body if isinstance(node, ast.FunctionDef)]
        namespace = dict(io=io, re=re, Document=Document)
        exec(compile(ast.Module(body=functions, type_ignores=[]), 'app.py', 'exec'), namespace)
        doc = Document(namespace['markdown_to_docx_buffer'](
            '1. 문제\n① 보기\n\n| 항목 | 값 |\n| -- | -- |\n| A | 10 |'))
        self.assertEqual(doc.paragraphs[0].text, '1. 문제\n① 보기')
        self.assertEqual(doc.tables[0].cell(1, 1).text, '10')

    def test_structure_and_math(self):
        doc = Document(transcription_to_docx_buffer(
            '# 제목\n\n첫 문단\n이어지는 줄\n\n둘째 **강조** 문단\n'
            '## 소제목\n- 항목\n  - 하위\n3. 셋째\n4. 넷째\n\n'
            '| 값 | 식 |\n| --- | ---: |\n| a\\|b<br>다음 | $x_i^2$ |\n\n'
            '$$\\frac{a_1}{b_2}$$\n[판독 불가]'
        ))
        self.assertEqual(doc.paragraphs[0].style.name, 'Heading 1')
        self.assertEqual(doc.paragraphs[1].text, '첫 문단\n이어지는 줄')
        self.assertTrue(any(r.bold and r.text == '강조' for r in doc.paragraphs[2].runs))
        self.assertEqual(doc.paragraphs[4].style.name, 'List Bullet')
        self.assertGreater(doc.paragraphs[5].paragraph_format.left_indent,
                           doc.paragraphs[4].paragraph_format.left_indent)
        self.assertEqual(doc.paragraphs[6].text, '3. 셋째')
        self.assertEqual(doc.tables[0].cell(1, 0).text, 'a|b\n다음')
        self.assertEqual(doc.tables[0].cell(1, 1).text, '$x_i^2$')
        self.assertIn('$$\\frac{a_1}{b_2}$$', doc.paragraphs[-1].text)

    def test_ragged_table_and_empty_document(self):
        doc = Document(transcription_to_docx_buffer('| A | B |\n| -- | -- |\n| C |'))
        self.assertEqual(doc.tables[0].cell(1, 1).text, '')
        self.assertEqual(len(Document(transcription_to_docx_buffer('')).tables), 0)


class AppTests(unittest.TestCase):
    def run_conversion(self, mode, pdf=False, response='원문 내용'):
        upload = io.BytesIO(b'%PDF-test')
        if not pdf:
            Image.new('RGB', (8, 8)).save(upload := io.BytesIO(), format='PNG')
        upload.name = 'sample.pdf' if pdf else 'sample.png'
        upload.type = 'application/pdf' if pdf else 'image/png'
        model = MagicMock()
        if isinstance(response, Exception):
            model.generate_content.side_effect = response
        else:
            model.generate_content.return_value.text = response
        with patch('streamlit.file_uploader', return_value=upload), \
                patch('google.generativeai.configure'), \
                patch('google.generativeai.GenerativeModel', return_value=model) as model_factory:
            app = AppTest.from_file(str(Path(__file__).with_name('app.py'))).run(timeout=30)
            self.assertFalse(app.exception)
            app.radio[0].set_value(mode)
            app.text_input[0].set_value('test-key').run()
            app.button[0].click().run()
            self.assertFalse(app.exception)
            model_factory.assert_called_once_with('gemini-3.8-flash')
        return app, model

    def test_both_modes_and_input_types(self):
        for mode in ['시험 문제 추출', '일반 내용 그대로 전사']:
            for pdf in [False, True]:
                with self.subTest(mode=mode, pdf=pdf):
                    app, model = self.run_conversion(mode, pdf)
                    content = model.generate_content.call_args.args[0]
                    self.assertEqual(content[0] == TRANSCRIPTION_PROMPT, mode == '일반 내용 그대로 전사')
                    self.assertEqual(isinstance(content[1], dict), pdf)
                    self.assertEqual(app.session_state['gemini_markdown'], '원문 내용')
                    self.assertTrue(app.session_state['download_file_name'].endswith('.docx'))
                    app.radio[0].set_value('시험 문제 추출' if mode != '시험 문제 추출' else '일반 내용 그대로 전사').run()
                    self.assertNotIn('gemini_markdown', app.session_state)

    def test_empty_response(self):
        app, _ = self.run_conversion('일반 내용 그대로 전사', response='  ')
        self.assertNotIn('gemini_markdown', app.session_state)
        self.assertTrue(app.warning)

    def test_api_error(self):
        app, _ = self.run_conversion('일반 내용 그대로 전사', response=RuntimeError('test failure'))
        self.assertNotIn('gemini_markdown', app.session_state)
        self.assertTrue(app.error)

    def test_default_and_disabled_button(self):
        app = AppTest.from_file(str(Path(__file__).with_name('app.py'))).run(timeout=30)
        self.assertEqual(app.radio[0].value, '시험 문제 추출')
        self.assertTrue(app.button[0].disabled)


if __name__ == '__main__':
    unittest.main()
